import json
import openpyxl
from openpyxl.utils import get_column_letter

# Load database dump
try:
    with open("extra/database_dump.json", "r", encoding="utf-8") as f:
        db = json.load(f)
except Exception as e:
    print(f"Error loading database_dump.json: {e}")
    exit(1)

kapans = db.get("kapans", [])
transfers = db.get("transfers", [])
rough_lots = db.get("roughLots", [])

# Sort kapans by createdDate so they are written in chronological order
kapans.sort(key=lambda x: x.get("createdDate", ""))

print(f"Loaded {len(kapans)} kapans, {len(transfers)} transfers, {len(rough_lots)} rough lots.")

# Load Excel template
wb_path = "Kaapan_Stock_Ledger_Template.xlsx"
try:
    wb = openpyxl.load_workbook(wb_path)
    sheet = wb["Detailed_Ledger"]
except Exception as e:
    print(f"Error opening Excel file: {e}")
    exit(1)

# List of departments in horizontal order
DEPTS_ORDER = [
    "Galaxy",
    "AP OK",
    "4P",
    "4P OK RT BAAKI",
    "RT",
    "RT OK KHATA BAAKI",
    "KHATA",
    "OK KAPAN (ઓકે કાપણ)"
]

# Mapping of department to its starting column index (1-based)
# AB: 28, AH: 34, AN: 40, AT: 46, AZ: 52, BF: 58, BL: 64, BR: 70
DEPT_COLUMNS = {
    "Galaxy": 28,
    "AP OK": 34,
    "4P": 40,
    "4P OK RT BAAKI": 46,
    "RT": 52,
    "RT OK KHATA BAAKI": 58,
    "KHATA": 64,
    "OK KAPAN (ઓકે કાપણ)": 70
}

# Clear old entries in Detailed_Ledger starting from row 5
# Only clear the input columns, leave the formulas untouched!
# Input columns to clear: A, B, C, D, E, F, H, I, J, M, N, P, R, S and all In/Out columns for departments
input_cols = [1, 2, 3, 4, 5, 6, 8, 9, 10, 13, 14, 16, 18, 19]
for d_col in DEPT_COLUMNS.values():
    input_cols.extend([d_col, d_col+1, d_col+2, d_col+3]) # In Pcs, In Cts, Out Pcs, Out Cts

print("Clearing old input cells in rows 5 to 100...")
for r in range(5, 101):
    for c in input_cols:
        sheet.cell(row=r, column=c).value = None

# Write kapans data
for idx, k in enumerate(kapans):
    r_idx = idx + 5  # Start writing at row 5
    if r_idx > 100:
        print("Warning: Exceeded 100 rows. Skipping remaining kapans.")
        break
        
    kapan_no = k.get("kapanNo")
    print(f"Writing Kapan {kapan_no} at row {r_idx}...")
    
    # 1. Look up Rough Lot details
    rough_id = k.get("roughId")
    rough = next((r for r in rough_lots if r.get("id") == rough_id), {})
    rough_name = rough.get("name", k.get("roughName", ""))
    rough_rate = rough.get("rate", "")
    
    # 2. Write basic columns
    sheet.cell(row=r_idx, column=1, value=idx + 1)      # A: No.
    sheet.cell(row=r_idx, column=2, value=rough_name)    # B: Rough No.
    sheet.cell(row=r_idx, column=3, value=k.get("vigat", "")) # C: Article Name (use Kapan comments)
    sheet.cell(row=r_idx, column=4, value=kapan_no)     # D: Kapan No.
    sheet.cell(row=r_idx, column=5, value=k.get("nang")) # E: Rough Pcs
    sheet.cell(row=r_idx, column=6, value=k.get("roughWeight", k.get("carat"))) # F: Rough Weight
    sheet.cell(row=r_idx, column=8, value=rough_rate)   # H: Rough Rate
    
    # 3. Write Planning & intermediate outputs
    sheet.cell(row=r_idx, column=9, value=k.get("makeablePiece")) # I: MK Pcs
    sheet.cell(row=r_idx, column=10, value=k.get("makeableVajan")) # J: Exp Wt
    sheet.cell(row=r_idx, column=13, value=k.get("fourPNang"))   # M: 4P Pcs
    sheet.cell(row=r_idx, column=14, value=k.get("fourPCt"))     # N: 4P Wt
    sheet.cell(row=r_idx, column=16, value=k.get("rtCt"))         # P: RT Wt
    
    # 4. If completed, write final Polish outputs
    if k.get("currentDept") == "OK KAPAN (ઓકે કાપણ)":
        sheet.cell(row=r_idx, column=18, value=k.get("carat"))   # R: Polish Wt
        sheet.cell(row=r_idx, column=19, value=k.get("nang"))    # S: Polish Pcs

    # 5. Track departmental transitions
    # Initialize department In/Out dict
    dept_data = {d: {"in_pcs": None, "in_cts": None, "out_pcs": None, "out_cts": None} for d in DEPTS_ORDER}
    
    # Set Galaxy In (start of Kapan)
    dept_data["Galaxy"]["in_pcs"] = k.get("nang")
    dept_data["Galaxy"]["in_cts"] = k.get("roughWeight", k.get("carat"))
    
    # Get transfers for this Kapan, sorted by timestamp
    k_transfers = [t for t in transfers if t.get("kapanNo") == kapan_no]
    k_transfers.sort(key=lambda x: x.get("timestamp", ""))
    
    for t in k_transfers:
        f_dept = t.get("fromDept")
        t_dept = t.get("toDept")
        pcs = t.get("nang")
        cts = t.get("carat")
        
        # Populate direct from/to dept values
        if f_dept in dept_data:
            dept_data[f_dept]["out_pcs"] = pcs
            dept_data[f_dept]["out_cts"] = cts
        if t_dept in dept_data:
            dept_data[t_dept]["in_pcs"] = pcs
            dept_data[t_dept]["in_cts"] = cts
            
        # Fill in skipped intermediate departments if any
        if f_dept in DEPTS_ORDER and t_dept in DEPTS_ORDER:
            idx_from = DEPTS_ORDER.index(f_dept)
            idx_to = DEPTS_ORDER.index(t_dept)
            if idx_to > idx_from + 1:
                # E.g. Galaxy -> 4P (skipped AP OK)
                for mid_idx in range(idx_from + 1, idx_to):
                    mid_dept = DEPTS_ORDER[mid_idx]
                    dept_data[mid_dept]["in_pcs"] = pcs
                    dept_data[mid_dept]["in_cts"] = cts
                    dept_data[mid_dept]["out_pcs"] = pcs
                    dept_data[mid_dept]["out_cts"] = cts

    # If Kapan is completed, the final department output equals its input
    if k.get("currentDept") == "OK KAPAN (ઓકે કાપણ)":
        dept_data["OK KAPAN (ઓકે કાપણ)"]["out_pcs"] = k.get("nang")
        dept_data["OK KAPAN (ઓકે કાપણ)"]["out_cts"] = k.get("carat")
        
    # Write departmental In/Out values into sheet
    for d, cols_idx in DEPT_COLUMNS.items():
        info = dept_data[d]
        sheet.cell(row=r_idx, column=cols_idx, value=info["in_pcs"])      # In Pcs
        sheet.cell(row=r_idx, column=cols_idx+1, value=info["in_cts"])    # In Cts
        sheet.cell(row=r_idx, column=cols_idx+2, value=info["out_pcs"])   # Out Pcs
        sheet.cell(row=r_idx, column=cols_idx+3, value=info["out_cts"])   # Out Cts

# Save workbook
try:
    wb.save(wb_path)
    print(f"Successfully uploaded Kapan details and saved workbook to: {wb_path}")
except PermissionError:
    print(f"ERROR: Permission Denied. Please close '{wb_path}' and run again.")
except Exception as e:
    print(f"Error saving workbook: {e}")
